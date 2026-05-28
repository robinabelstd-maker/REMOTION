"""
Fallback: export progress.json to an Excel (.xlsx) file.
Useful if you don't have Google Sheets set up yet.
Output: Desktop/salesdock_export/transactions.xlsx
"""

import json
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DESKTOP = Path(os.path.expanduser("~")) / "Desktop"
EXPORT_DIR = DESKTOP / "salesdock_export"
PROGRESS_FILE = EXPORT_DIR / "progress.json"
OUTPUT_FILE = EXPORT_DIR / "transactions.xlsx"

COLUMN_ORDER = [
    "Transaction ID",
    "Customer name",
    "Delivery address",
    "Correspondence address",
    "Email",
    "Phone",
    "Date of birth",
    "IBAN",
    "IBAN account name",
    "Product name",
    "Status",
    "Seller",
    "Organisation",
    "Flow",
    "Date created",
    "Date updated",
    "Contract sign date",
    "Outgoing ID",
    "Outgoing date",
    "Sale type",
    "Type",
    "Monthly price",
]


def export():
    if not PROGRESS_FILE.exists():
        print(f"ERROR: {PROGRESS_FILE} not found. Run scraper.py first.")
        return

    with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    transactions = data.get("transactions", [])
    print(f"Exporting {len(transactions)} transactions to Excel …")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, record in enumerate(transactions, start=2):
        for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
            ws.cell(row=row_idx, column=col_idx, value=record.get(col_name, ""))

    # Auto-fit column widths (approximate)
    for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row_idx in range(2, min(len(transactions) + 2, 200)):
            val = ws.cell(row=row_idx, column=col_idx).value or ""
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)
    print(f"Saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    export()
